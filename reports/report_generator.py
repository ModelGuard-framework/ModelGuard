
from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document

from core.validator import overall_status


def generate_report(
    results: dict,
    output_path: str | Path = "reports/test_report.docx",
    model_path: str | Path | None = None,
    dataset_path: str | Path | None = None,
    label_column: str | None = None,
    template_path: str | Path | None = None,
):
    """
    Generate the ModelGuard Word report.

    If template_path points to a Word document, it is used as the base
    document. Otherwise a standard ModelGuard report layout is created.
    """
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    if template_path and Path(template_path).exists():
        document = Document(str(template_path))
    else:
        document = Document()

    document.add_heading("MODEL GUARD — ML TEST REPORT", 0)
    document.add_paragraph(
        f"Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    )

    document.add_heading("Configuration", level=1)
    if model_path:
        document.add_paragraph(f"Model: {Path(model_path).name}")
    if dataset_path:
        document.add_paragraph(f"Dataset: {Path(dataset_path).name}")
    document.add_paragraph(f"Label column: {label_column or 'Not provided'}")

    document.add_heading("Testing Results", level=1)

    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    header = table.rows[0].cells
    header[0].text = "Testing Stage"
    header[1].text = "Status"
    header[2].text = "Details"

    for stage, item in results.items():
        row = table.add_row().cells
        row[0].text = stage
        row[1].text = item.get("status", "NOT RUN")
        row[2].text = " ".join(item.get("details", []))

    document.add_heading("Metrics", level=1)
    has_metrics = False
    for stage, item in results.items():
        metrics = item.get("metrics") or {}
        if metrics:
            has_metrics = True
            document.add_paragraph(stage)
            for name, value in metrics.items():
                document.add_paragraph(
                    f"{name.replace('_', ' ').title()}: {value:.4f}"
                    if isinstance(value, float)
                    else f"{name.replace('_', ' ').title()}: {value}",
                    style="List Bullet",
                )

    if not has_metrics:
        document.add_paragraph("No ground-truth labels were supplied, so performance metrics were skipped.")

    final_status = overall_status(results)
    document.add_heading("Overall Result", level=1)
    document.add_paragraph(f"Overall Core Status: {final_status}")
    document.add_paragraph(
        "Scope note: 'Overall Core Status' above covers only the 11 backend "
        "validation stages listed in this report. It does not include the "
        "separate Selenium UI Validation result. The ModelGuard Dashboard's "
        "'Overall Status' metric combines both, so it can show FAIL even when "
        "this report shows PASS, if the automated browser (Selenium) test failed "
        "or has not completed. Check the Validation page's 'Selenium UI "
        "Validation' section for that result."
    )

    document.add_paragraph(
        "Note: Data Poisoning Detection reports basic indicators of suspicious data "
        "patterns. It does not prove that a dataset is free from poisoning."
    )

    document.save(output_path)
    return output_path
